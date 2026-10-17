"""Day 31 - Cash Flow Intelligence module.

Per company this computes:

* CFO quality  - average of CFO / PAT over the last 5 years
                 (years with non-positive PAT are excluded as the ratio is not meaningful)
* CapEx intensity - |investing activity| / sales x 100 (latest year)
* FCF CAGR (5Y) and FCF conversion (FCF / PAT, latest year)
* Distress signal - CFO < 0 AND CFF > 0 in the latest year
* Deleveraging flag - CFF < 0 AND borrowings lower than the previous year
"""
from __future__ import annotations

from typing import Dict, Optional

import numpy as np
import pandas as pd
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from config import settings
from src.analytics import capital_allocation
from src.utils.data_loader import DataStore
from src.utils.helpers import NAN, cagr_pct, is_nan, safe_div
from src.utils.logger import get_logger

log = get_logger(__name__)

SPEC_COLUMNS = [
    "company_id", "sector", "cfo_quality_score", "cfo_quality_label", "capex_intensity_pct",
    "capex_label", "fcf_cagr_5yr", "fcf_conversion_pct", "distress_flag", "deleveraging_flag",
    "capital_allocation_label",
]


def cfo_quality_label(score: float) -> str:
    if is_nan(score):
        return "Insufficient Data"
    if score > settings.CFO_QUALITY_HIGH:
        return "High Quality"
    if score >= settings.CFO_QUALITY_MODERATE:
        return "Moderate"
    return "Accrual Risk"


def capex_label(pct: float) -> str:
    if is_nan(pct):
        return "Insufficient Data"
    if pct < settings.CAPEX_ASSET_LIGHT:
        return "Asset Light"
    if pct <= settings.CAPEX_CAPITAL_INTENSIVE:
        return "Moderate"
    return "Capital Intensive"


def _company_row(comp: pd.Series, frame: pd.DataFrame, allocation: Dict) -> Dict:
    row = {"company_id": comp["company_id"], "sector": comp["sector"],
           "ticker": comp["ticker"], "company_name": comp["company_name"]}
    if frame.empty:
        row.update(cfo_quality_score=NAN, capex_intensity_pct=NAN, fcf_cagr_5yr=NAN,
                   fcf_conversion_pct=NAN, distress_flag=False, deleveraging_flag=False,
                   latest_cfo=NAN, latest_cff=NAN, latest_net_profit=NAN, latest_year=NAN)
    else:
        f = frame.sort_values("year").set_index("year")
        ly = int(f.index.max())
        window = f.loc[ly - settings.CFO_QUALITY_WINDOW_YEARS + 1: ly]
        ratios = (window["operating_activity"] / window["net_profit"]).where(window["net_profit"] > 0)
        score = float(ratios.mean()) if ratios.notna().any() else NAN

        cfo, cfi, cff = (f["operating_activity"].get(ly, NAN), f["investing_activity"].get(ly, NAN),
                         f["financing_activity"].get(ly, NAN))
        sales, pat = f["sales"].get(ly, NAN), f["net_profit"].get(ly, NAN)
        fcf = f["fcf"].get(ly, NAN)
        borrow_now, borrow_prev = f["borrowings"].get(ly, NAN), f["borrowings"].get(ly - 1, NAN)

        capex = safe_div(abs(cfi), sales) * 100 if not is_nan(cfi) else NAN
        conversion = safe_div(fcf, pat) * 100 if not is_nan(pat) and pat > 0 else NAN
        row.update(
            cfo_quality_score=round(score, 3) if not is_nan(score) else NAN,
            capex_intensity_pct=round(capex, 2) if not is_nan(capex) else NAN,
            fcf_cagr_5yr=round(cagr_pct(f["fcf"].get(ly - 5, NAN), fcf, 5), 2),
            fcf_conversion_pct=round(conversion, 2) if not is_nan(conversion) else NAN,
            distress_flag=bool(not is_nan(cfo) and not is_nan(cff) and cfo < 0 and cff > 0),
            deleveraging_flag=bool(not is_nan(cff) and not is_nan(borrow_now) and not is_nan(borrow_prev)
                                   and cff < 0 and borrow_now < borrow_prev),
            latest_cfo=cfo, latest_cff=cff, latest_net_profit=pat, latest_year=ly)
    row["cfo_quality_label"] = cfo_quality_label(row["cfo_quality_score"])
    row["capex_label"] = capex_label(row["capex_intensity_pct"])
    row["capital_allocation_label"] = allocation.get(comp["company_id"], "Not Available")
    return row


def compute_cashflow_kpis(store: DataStore, allocation: Optional[Dict] = None) -> pd.DataFrame:
    """Full per-company table (spec columns + helper columns used by the distress report)."""
    allocation = allocation or {}
    rows = [_company_row(comp, store.frame_for(comp["company_id"]), allocation)
            for _, comp in store.companies.iterrows()]
    return pd.DataFrame(rows)


def distress_alerts(full: pd.DataFrame) -> pd.DataFrame:
    d = full[full["distress_flag"]].copy()
    d = d.rename(columns={"latest_cfo": "cfo_value", "latest_cff": "cff_value"})
    cols = ["company_id", "ticker", "company_name", "sector", "latest_year",
            "cfo_value", "cff_value", "latest_net_profit"]
    return d[cols].sort_values("cfo_value").reset_index(drop=True)


# --------------------------------------------------------------------------- #
# Excel export
# --------------------------------------------------------------------------- #
_FONT = "Arial"
_HEADER_FILL = PatternFill("solid", fgColor="0B1F3A")
_RED = PatternFill("solid", fgColor="F8D7DA")
_GREEN = PatternFill("solid", fgColor="D4EDDA")
_AMBER = PatternFill("solid", fgColor="FFF3CD")
_THIN = Side(style="thin", color="D0D5DD")

_LABEL_FILLS = {
    "High Quality": _GREEN, "Moderate": _AMBER, "Accrual Risk": _RED,
    "Asset Light": _GREEN, "Capital Intensive": _AMBER,
    "Distress Signal": _RED,
}


def _style_sheet(ws, widths: Dict[str, float], number_formats: Optional[Dict[str, str]] = None) -> None:
    for cell in ws[1]:
        cell.font = Font(name=_FONT, bold=True, color="FFFFFF", size=10)
        cell.fill = _HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[1].height = 30
    headers = [c.value for c in ws[1]]
    for idx, name in enumerate(headers, start=1):
        ws.column_dimensions[get_column_letter(idx)].width = widths.get(name, 16)
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.font = Font(name=_FONT, size=10)
            cell.border = Border(bottom=_THIN)
            header = headers[cell.column - 1]
            if number_formats and header in number_formats:
                cell.number_format = number_formats[header]
            if isinstance(cell.value, str) and cell.value in _LABEL_FILLS:
                cell.fill = _LABEL_FILLS[cell.value]
            if header in ("distress_flag", "deleveraging_flag") and cell.value in (True, "TRUE"):
                cell.fill = _RED if header == "distress_flag" else _GREEN
            if header in ("distress_flag", "deleveraging_flag"):
                cell.alignment = Alignment(horizontal="center")
    ws.freeze_panes = "B2"
    ws.auto_filter.ref = ws.dimensions


NOTES = [
    ("Sheet", "Description"),
    ("Cash Flow Intelligence", "One row per company with the 11 sprint-spec columns."),
    ("Distribution", "Count of companies per CFO-quality, CapEx and capital-allocation label."),
    ("CFO quality score", f"Average of CFO / PAT over the last {settings.CFO_QUALITY_WINDOW_YEARS} years; "
                          "years with PAT <= 0 are excluded. >1.0 High Quality, 0.5-1.0 Moderate, <0.5 Accrual Risk."),
    ("CapEx intensity", "abs(investing activity) / sales x 100, latest year. <3% Asset Light, 3-8% Moderate, >8% Capital Intensive."),
    ("FCF CAGR / conversion", "FCF CAGR over 5 years (blank unless both end points are positive); conversion = FCF / PAT x 100."),
    ("Distress flag", "Latest year CFO < 0 and CFF > 0 (financing is funding operating losses)."),
    ("Deleveraging flag", "Latest year CFF < 0 and borrowings lower than the previous year."),
]


def export_workbook(full: pd.DataFrame, path=settings.CASHFLOW_XLSX) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    main = full[SPEC_COLUMNS].copy()
    distribution = pd.concat([
        main["cfo_quality_label"].value_counts().rename_axis("label").reset_index(name="companies").assign(dimension="CFO quality"),
        main["capex_label"].value_counts().rename_axis("label").reset_index(name="companies").assign(dimension="CapEx intensity"),
        main["capital_allocation_label"].value_counts().rename_axis("label").reset_index(name="companies").assign(dimension="Capital allocation"),
    ])[["dimension", "label", "companies"]]

    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        main.to_excel(writer, sheet_name="Cash Flow Intelligence", index=False)
        distribution.to_excel(writer, sheet_name="Distribution", index=False)
        pd.DataFrame(NOTES[1:], columns=NOTES[0]).to_excel(writer, sheet_name="Notes", index=False)

        _style_sheet(writer.sheets["Cash Flow Intelligence"],
                     widths={"company_id": 12, "sector": 24, "cfo_quality_label": 20, "capex_label": 20,
                             "capital_allocation_label": 26, "cfo_quality_score": 16},
                     number_formats={"cfo_quality_score": "0.00", "capex_intensity_pct": "0.0",
                                     "fcf_cagr_5yr": "0.0", "fcf_conversion_pct": "0.0"})
        _style_sheet(writer.sheets["Distribution"], widths={"dimension": 22, "label": 28, "companies": 12})
        _style_sheet(writer.sheets["Notes"], widths={"Sheet": 26, "Description": 110})
        for row in writer.sheets["Notes"].iter_rows(min_row=2):
            row[1].alignment = Alignment(wrap_text=True, vertical="top")


def run(store: DataStore) -> pd.DataFrame:
    """Build cashflow_intelligence.xlsx (+ capital allocation labels) and distress_alerts.csv."""
    settings.OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    try:
        allocation = capital_allocation.latest_labels(store.table("capital_allocation"))
    except Exception as exc:  # file missing -> still produce the workbook
        log.warning("Capital allocation labels unavailable (%s); using 'Not Available'", exc)
        allocation = {}
    full = compute_cashflow_kpis(store, allocation)
    export_workbook(full)
    alerts = distress_alerts(full)
    alerts.to_csv(settings.DISTRESS_ALERTS_CSV, index=False)
    log.info("cashflow_intelligence.xlsx written (%d rows) | %d distress alerts", len(full), len(alerts))
    return full
